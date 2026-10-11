"""Reading files and writing the Excel report.

Test files are written with write_bytes: write_text() turns \\n into \\r\\n on
Windows, which would change the very line endings these tests check.
"""

import os

import pytest
from openpyxl import load_workbook

from ezcompare.engine import CompareOptions, compare
from ezcompare.export import export_excel
from ezcompare.textfile import CompareError, read_text_file, split_lines, text_from_string


def test_reads_utf8_and_reports_line_endings(tmp_path):
    path = tmp_path / "a.txt"
    path.write_bytes("one\r\ntwo\r\nthree\r\n".encode("utf-8"))
    source = read_text_file(str(path))
    assert source.lines == ("one", "two", "three")
    assert source.line_ending == "Windows (CRLF)"
    assert source.encoding == "UTF-8"


def test_reads_utf16_with_bom(tmp_path):
    path = tmp_path / "ps.log"
    path.write_bytes("PS output\nline 2\n".encode("utf-16"))
    source = read_text_file(str(path))
    assert source.lines == ("PS output", "line 2")
    assert source.encoding == "UTF-16"


def test_reads_windows_1252(tmp_path):
    path = tmp_path / "old.cfg"
    path.write_bytes("caf\xe9\n".encode("cp1252"))
    assert read_text_file(str(path)).lines == ("café",)


def test_line_endings_alone_are_not_differences(tmp_path):
    a, b = tmp_path / "a.txt", tmp_path / "b.txt"
    a.write_bytes(b"x\r\ny\r\n")
    b.write_bytes(b"x\ny")
    assert compare(read_text_file(str(a)), read_text_file(str(b))).identical


def test_form_feed_does_not_split_a_line():
    assert split_lines("a\x0cb\nc") == ("a\x0cb", "c")
    assert split_lines("") == ()
    assert split_lines("a\n") == ("a",)


def test_binary_file_is_refused(tmp_path):
    path = tmp_path / "firmware.bin"
    path.write_bytes(bytes(range(256)) * 4)
    with pytest.raises(CompareError, match="not a text file"):
        read_text_file(str(path))


def test_missing_file_is_explained(tmp_path):
    with pytest.raises(CompareError, match="can't be found"):
        read_text_file(str(tmp_path / "nope.txt"))


def test_folder_is_refused(tmp_path):
    with pytest.raises(CompareError, match="is a folder"):
        read_text_file(str(tmp_path))


def test_excel_report(tmp_path):
    left = text_from_string("hostname R1\nip 10.0.0.1\nshutdown\n=SUM(A1)\n")
    right = text_from_string("hostname R1\nip 10.0.0.2\nlogging on\n")
    options = CompareOptions()
    result = compare(left, right, options)
    out = tmp_path / "report.xlsx"
    rows = export_excel(result, left, right, options, str(out))
    assert rows == result.removed + result.added + result.changed
    wb = load_workbook(out)
    sheet = wb["Differences"]
    values = list(sheet.iter_rows(min_row=2, values_only=True))
    assert ("Changed", 2, "ip 10.0.0.1", 2, "ip 10.0.0.2") in values
    # A line that looks like a formula is stored as text, never run.
    cell = next(c for row in sheet.iter_rows(min_row=2) for c in row if c.value == "=SUM(A1)")
    assert cell.data_type == "s"
    assert "Summary" in wb.sheetnames
    # Nothing is left behind next to the report.
    assert [p.name for p in tmp_path.iterdir()] == ["report.xlsx"]


def test_report_never_overwrites_a_compared_file(tmp_path):
    path = tmp_path / "a.txt"
    path.write_bytes(b"x\n")
    left = read_text_file(str(path))
    right = text_from_string("y\n")
    result = compare(left, right)
    with pytest.raises(CompareError, match="one of the files being compared"):
        export_excel(result, left, right, CompareOptions(), str(path))
    assert path.read_bytes() == b"x\n"


def test_report_to_a_missing_folder_is_explained(tmp_path):
    left, right = text_from_string("a\n"), text_from_string("b\n")
    result = compare(left, right)
    with pytest.raises(CompareError, match="Can't save"):
        export_excel(result, left, right, CompareOptions(), os.path.join(str(tmp_path), "no", "r.xlsx"))
