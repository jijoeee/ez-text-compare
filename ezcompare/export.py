"""Save a comparison as an Excel workbook: the differences, and a summary.

No UI imports. Saved to a temp file in the destination folder and swapped in,
so a failure never leaves half a workbook behind.
"""

from __future__ import annotations

import os
import tempfile
from datetime import datetime
from typing import Iterator, Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import APP_NAME, __version__
from .engine import ADDED, CHANGED, MOVED, REMOVED, CompareOptions, CompareResult, Line, summary_lines
from .textfile import CompareError, TextSource

EXCEL_CELL_LIMIT = 32_767      # characters Excel accepts in one cell
EXCEL_ROW_LIMIT = 1_048_576

STATUS_LABELS = {
    REMOVED: "Only in left",
    ADDED: "Only in right",
    CHANGED: "Changed",
    MOVED: "Moved",
}
# Soft fills matching the app's light-mode line colours.
STATUS_FILLS = {
    REMOVED: "FBE8E8",
    ADDED: "E3F4E8",
    CHANGED: "FCF2D6",
    MOVED: "E3EDF9",
}


def _cell_text(text: str) -> str:
    """Text Excel will store as plain text: no control characters, not too long."""
    text = ILLEGAL_CHARACTERS_RE.sub(" ", text)
    if len(text) > EXCEL_CELL_LIMIT:
        text = text[:EXCEL_CELL_LIMIT - 1] + "…"
    return text


def difference_rows(result: CompareResult) -> Iterator[tuple[str, Optional[int], str, Optional[int], str]]:
    """(status, left line, left text, right line, right text) for every difference.

    A changed or moved pair is one row with both sides, so the report reads like the screen.
    """
    def side(ln: Optional[Line]) -> tuple[Optional[int], str]:
        if ln is None or ln.number is None:
            return None, ""
        return ln.number, ln.text

    if result.aligned:
        for left, right in zip(result.left, result.right):
            if left.kind == CHANGED:
                yield (CHANGED, *side(left), *side(right))
            elif left.kind == MOVED and left.partner is not None:
                # One row per move, written where the line left; its new place is on the right.
                yield (MOVED, *side(left), *side(result.right[left.partner]))
            elif left.kind == REMOVED:
                yield (REMOVED, *side(left), None, "")
            if right.kind == ADDED:
                yield (ADDED, None, "", *side(right))
        return

    for left in result.left:
        if left.kind == CHANGED and left.partner is not None:
            yield (CHANGED, *side(left), *side(result.right[left.partner]))
        elif left.kind == REMOVED:
            yield (REMOVED, *side(left), None, "")
    for right in result.right:
        if right.kind == ADDED:
            yield (ADDED, None, "", *side(right))


def _options_text(options: CompareOptions) -> list[tuple[str, str]]:
    ignores = [label for on, label in (
        (options.ignore_empty, "empty lines"),
        (options.ignore_all_spaces, "all spaces"),
        (options.ignore_changed_spaces and not options.ignore_all_spaces, "changes in spaces"),
        (options.ignore_case, "upper and lower case"),
        (options.ignore_timestamps, "dates and times"),
    ) if on]
    if options.ignore_pattern.strip():
        ignores.append(f"text matching  {options.ignore_pattern}")
    only = options.only_containing.strip()
    return [
        ("Compared", f"only lines containing \"{only}\" (any case)" if only else "every line"),
        ("Line order", "ignored (lines compared wherever they are)" if options.ignore_order
         else "matters (side-by-side diff)"),
        ("Ignored", ", ".join(ignores) if ignores else "nothing"),
    ]


def _describe(source: TextSource) -> str:
    if source.pasted:
        return f"{source.name} ({len(source.lines):,} lines)"
    return f"{source.path} ({len(source.lines):,} lines)"


def export_excel(result: CompareResult, left: TextSource, right: TextSource,
                 options: CompareOptions, path: str) -> int:
    """Write the workbook. Returns the number of difference rows written."""
    name = os.path.basename(path)
    folder = os.path.dirname(os.path.abspath(path)) or "."
    for source in (left, right):
        if source.path and os.path.abspath(source.path) == os.path.abspath(path):
            raise CompareError(f"{name} is one of the files being compared. Save the report "
                               "with another name.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Differences"
    header = ["Status", "Left line", f"Left: {left.name}", "Right line", f"Right: {right.name}"]
    ws.append([_cell_text(h) for h in header])
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold
    rows = 0
    for status, ln, lt, rn, rt in difference_rows(result):
        if rows >= EXCEL_ROW_LIMIT - 2:
            break
        ws.append([STATUS_LABELS[status], ln, _cell_text(lt), rn, _cell_text(rt)])
        rows += 1
        excel_row = rows + 1
        fill = PatternFill("solid", fgColor=STATUS_FILLS[status])
        for col in range(1, 6):
            ws.cell(excel_row, col).fill = fill
        # Lines that start with "=" are stored as text, never run as formulas.
        for col in (3, 5):
            cell = ws.cell(excel_row, col)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    for col, width in zip(range(1, 6), (14, 10, 70, 10, 70)):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{rows + 1}"
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top")

    info = wb.create_sheet("Summary")
    summary = [
        ("Report", f"{APP_NAME} {__version__}"),
        ("Created", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Left", _describe(left)),
        ("Right", _describe(right)),
        *_options_text(options),
        ("Result", "The files are the same" if result.identical
         else f"{result.differences:,} lines differ"),
    ]
    summary += [("", line) for line in summary_lines(result)]
    if rows < sum(1 for _ in difference_rows(result)):
        summary.append(("Note", f"Only the first {rows:,} differences fit in one Excel sheet."))
    for label, value in summary:
        info.append([label, _cell_text(value)])
    for cell in info["A"]:
        cell.font = bold
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 100

    try:
        fd, tmp = tempfile.mkstemp(prefix=".ezcompare-", suffix=".tmp", dir=folder)
        os.close(fd)
    except OSError as exc:
        raise CompareError(f"Can't save {name} in that folder: {exc.strerror or 'disk error'}. "
                           "Choose another folder.") from None
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        _silent_remove(tmp)
        raise CompareError(f"Can't save {name}. If it is open in Excel, close it and try again, "
                           "or save with another name.") from None
    except OSError as exc:
        _silent_remove(tmp)
        raise CompareError(f"Can't save {name}: {exc.strerror or 'disk error'}.") from None
    except BaseException:
        _silent_remove(tmp)
        raise
    return rows


def _silent_remove(path: str) -> None:
    try:
        os.remove(path)
    except OSError:
        pass
